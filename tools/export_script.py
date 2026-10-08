#!/usr/bin/env python3
"""Export the current, resolved campaign script without rewriting dialogue.

Markdown requires only Python and Node. Add --docx for an editable Word copy
using python-docx. --check compares exports with the actual authoring source.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[1]
TITLE = "THE CRITIC: COMING ATTRACTIONS"
SOURCE_FILES = ("data/cutscenes.js", "data/campaign.js", "engine.js", "app.js", "config.js")
ROUTES = {"hero": "Jay route", "franklin": "Franklin route"}
ACTOR_NAMES = {
    "hero": "Jay", "selected": "selected player", "franklin": "Franklin",
    "duke": "Duke", "marty": "Marty", "pizzeria-boss": "Violent Austrian Rabbi",
    "spike": "Spike", "sherm-punch": "Shermometer v1", "sherm-shove": "Shermometer v2",
    "sherm-slam": "Shermometer v3", "striped": "Fred K", "raptor": "JP Raptor Esq",
    "bear": "Accordion Bear", "hippo": "Green Hippo",
    "broadcast-rig": "Duke’s Broadcast System",
}
ENVIRONMENTS = {"studio": "Coming Attractions studio", "broadway": "Broadway",
                "subway": "Last Train Uptown", "rooftop": "Above the Avenue",
                "theater": "Theater District", "cinema": "Palace Cinema",
                "pizzeria": "Little Italy pizzeria", "broadcast": "Broadcast Tower"}
UI_IDS = {"rosterNote": "Character selection", "rewardHeading": "Results heading",
          "rewardText": "Results message", "unlockStatus": "Boss and unlock status"}
SCENE_ORDER = ["opening", "stage-02-intro", "stage-03-intro", "stage-04-intro",
               "stage4-clear", "stage-05-intro", "boss-cinema-intro",
               "boss-projection-defeat", "boss-cinema-defeat", "stage-06-intro",
               "boss-spike-intro", "boss-spike-defeat", "stage-07-intro",
               "boss-broadcast-intro", "boss-broadcast-defeat", "boss-duke-intro",
               "boss-duke-defeat", "ending"]


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def node_binary() -> str:
    return os.environ.get("CODEX_PRIMARY_RUNTIME_NODE", "node")


def source_model() -> dict:
    script = r"""
const vm=require('vm'),fs=require('fs');
const d=require('./data/cutscenes.js'),c=require('./data/campaign.js');
const sandbox={};sandbox.window=sandbox;vm.runInNewContext(fs.readFileSync('config.js','utf8'),sandbox);
const config=sandbox.window.BRAWLER_CONFIG;
const stages=c.STAGES||globalThis.CriticCampaign?.STAGES||[];
const scenes=Object.values(d.scenes).map(scene=>({raw:scene,routes:Object.fromEntries(
 ['hero','franklin'].map(route=>[route,d.resolveScene(scene,route)]))}));
const ui=[],app=fs.readFileSync('app.js','utf8');
const ids=['rosterNote','rewardHeading','rewardText','unlockStatus'];
for(const id of ids){const re=new RegExp("\\$\\('"+id+"'\\)\\.textContent=([^;]+);",'g');
 for(const match of app.matchAll(re)){
  const seen=new Set();
  for(const playerKind of ['hero','franklin'])for(const franklinUnlocked of [true,false])
  for(const stage of [3,6])for(const stage4Eligible of [true,false])
  for(const circuitFight of [false,true])for(const off of [0,1,2,3]){
   const game={playerKind,stage,stage4Eligible,deathsByStage:[0,0,0,stage4Eligible?0:1]},
    profile={franklinUnlocked},replay=playerKind==='franklin';
   const text=vm.runInNewContext(match[1],{game,profile,replay,circuitFight,off});
   if(typeof text==='string'&&text&&!seen.has(text)){seen.add(text);ui.push({id,text});}
  }
 }}
for(const match of app.matchAll(/toast\((['"])([^'"]*(?:UNLOCKED|Stage 4)[^'"]*)\1\)/g))
 ui.push({id:'toast',text:match[2]});
console.log(JSON.stringify({version:config.version,scenes,stages,ui}));
"""
    raw = subprocess.check_output([node_binary(), "-e", script], cwd=ROOT, text=True)
    model = json.loads(raw)
    model["sourceDigest"] = sha(json.dumps(model, ensure_ascii=False, sort_keys=True).encode())
    model["sourceHashes"] = {name: sha((ROOT / name).read_bytes()) for name in SOURCE_FILES}
    return model


def trigger(scene_id: str, model: dict) -> str:
    if scene_id == "opening":
        return "New Game, before Stage 1. Continue resumes the saved checkpoint."
    if scene_id == "ending":
        return "After Duke’s defeat settles. Marty is released before results."
    if scene_id == "stage4-clear":
        return "After Stage 4 is cleared. The Franklin line appears only on Jay’s route."
    if scene_id == "boss-projection-intro":
        return "When the player reaches the Palace movie-screen encounter area."
    if scene_id == "boss-projection-defeat":
        return "When all three projection circuits have been disabled. The cinema boss must also be defeated to leave."
    for number, stage in enumerate(model["stages"], 1):
        if stage.get("intro") == scene_id:
            return f"On entry to Stage {number}, {stage['name']}."
        for key in ("boss", "finalBoss"):
            boss = stage.get(key) or {}
            if boss.get("intro") == scene_id:
                return f"Before the {ACTOR_NAMES.get(boss['kind'], boss['kind'])} fight in Stage {number}, {stage['name']}."
    return {
        "boss-cinema-defeat": "After the Violent Austrian Rabbi is defeated. The circuits must also be disabled to leave.",
        "boss-spike-defeat": "After Spike is defeated, before the exit toward the Broadcast Tower.",
        "boss-broadcast-defeat": "After the machine is destroyed. Enemy transmissions stop; Marty remains confined and Duke’s fight follows.",
        "boss-duke-defeat": "After Duke is defeated. His defeat permits the cage to open and the rescue to proceed.",
        "boss-pizzeria-defeat": "After the pizzeria boss is defeated.",
    }.get(scene_id, "At the corresponding authored story event.")


def human_action(shot: dict, route: str) -> str:
    if shot.get("gameplayEntry"):
        return "The scene ends and the selected player drops onto the actual Broadway gameplay canvas, landing through the ordinary jump and recovery physics."
    notes = []
    actors = [a for a in shot.get("actors", []) if not a.get("routes") or route in a["routes"]]
    for a in actors:
        kind = a.get("character", a["id"])
        name = "Selected player" if kind == "selected" else ACTOR_NAMES.get(kind, kind)
        animation = a.get("animation", "idle").replace("-", " ")
        if a.get("image") and "seated" in a["image"]:
            animation = "seated in the review chair"
        motion = a.get("motion")
        if motion:
            dx = motion.get("toX", a.get("x", 0)) - motion.get("fromX", a.get("x", 0))
            dy = motion.get("toY", a.get("y", 0)) - motion.get("fromY", a.get("y", 0))
            direction = "right" if dx > 0 else "left" if dx < 0 else "in place"
            if dy:
                direction += " and down" if dy > 0 else " and up"
            if a.get("emerging"):
                notes.append(f"{name} emerges from the cinema screen and moves {direction} onto the combat plane")
            else:
                notes.append(f"{name}: {animation}, moving {direction} over {motion.get('duration', 1.4):g}s")
        else:
            notes.append(f"{name}: {animation}")
        if a.get("afterAnimation"):
            notes[-1] += f", then {a['afterAnimation'].replace('-', ' ')} after {a.get('after', 0):g}s"
    if shot.get("cage"):
        c = shot["cage"]
        notes.append("Marty’s cage is open" if c.get("open") else "Duke transports Marty’s closed cage" if c.get("carried") else "Marty is confined in the cage")
    if shot.get("curtainReveal"):
        notes.append("the curtain reveals Marty for the first time")
    if shot.get("powered") is not None:
        notes.append("broadcast monitors are powered" if shot["powered"] else "the broadcast monitors lose power")
    if shot.get("emissions"):
        cast = ", ".join(ACTOR_NAMES.get(e["character"], e["character"]) for e in shot["emissions"])
        if shot.get("emissionGroup") and any(e.get("approach") for e in shot["emissions"]):
            if shot.get("id") == "screen-emergence":
                notes.append(f"screens release {cast} in succession; they remain visible and creep toward Jay, clear of Marty’s cage")
            else:
                notes.append("the same screen-born enemies remain visible and advance toward Jay")
            actions = shot.get("emissionActions", {})
            if actions.get("attraction-0", {}).get("afterAnimation") == "attack":
                notes.append("Shermometer v1 attacks Jay as the other enemies approach")
        else:
            notes.append(f"screens release {cast} in succession")
    booth = shot.get("booth")
    if booth:
        notes.append({"shadow": "eyes move in the dark projection booth", "lit": "the booth light reveals the projectionist", "off": "the projection booths go dark"}.get(booth["phase"], "projection booth active"))
    if shot.get("screenForeshadow"):
        notes.append("the Violent Austrian Rabbi is foreshadowed inside the cinema screen")
    if shot.get("door"):
        notes.append("the pizzeria door opens" if shot["door"].get("opening") else "the pizzeria door is open")
    for key, desc in (("brokenWindow", "the studio window breaks"), ("glass", "glass scatters"),
                      ("alarm", "alarms pulse"), ("flash", "a white flash punctuates the beat"),
                      ("shake", "the scene shakes")):
        if shot.get(key):
            notes.append(desc)
    if shot.get("sound"):
        notes.append("sound cue: " + shot["sound"])
    return "; ".join(notes) + ("." if notes else "")


def timing(shot: dict) -> str:
    if shot.get("gameplayEntry"):
        return "Immediate transition to Stage 1 gameplay; no cutscene hold."
    result = f"Automatic advance after {shot['auto']:g}s" if shot.get("auto") else "Player advances"
    if shot.get("minTime"):
        result += f"; earliest advance {shot['minTime']:g}s"
    return result + "."


def blocks(model: dict) -> list[tuple[str, str]]:
    out = [("title", TITLE), ("subtitle", f"Full campaign script v{model['version']}"),
           ("body", "Edit the dialogue under each speaker. Scene and beat identifiers connect every line to its place in the game. Shared exchanges appear once; Jay and Franklin alternatives are shown together."),
           ("body", "The opening follows Jay’s confrontation with Duke on both routes. Franklin is present as an ally on his route. Jay remains Marty’s father throughout the story."),
           ("heading1", "Campaign order")]
    for number, stage in enumerate(model["stages"], 1):
        purpose = stage["purpose"]
        if stage.get("boss", {}).get("alternateKind"):
            purpose += " Franklin’s route faces Shermometer v3 here."
        out.append(("body", f"{number}. {stage['name']}. {purpose}"))
    out.append(("body", "Final sequence: broadcast system defeated, Duke fight, Duke defeated, Marty rescued, selected player reaction, results."))
    scenes = sorted(model["scenes"],
                    key=lambda s: SCENE_ORDER.index(s["raw"]["id"]) if s["raw"]["id"] in SCENE_ORDER else len(SCENE_ORDER))
    optional = []
    for scene in scenes + optional:
        raw = scene["raw"]
        if raw["id"] == "boss-projection-intro":
            out.append(("heading1", "Projection booth encounter introduction"))
        else:
            out.append(("opening" if raw["id"] == "opening" else "heading1", raw["title"].replace(" / ", " ").replace("’", "").replace("!", "")))
        out.append(("meta", "Scene " + raw["id"] + " | Onscreen title: " + raw["title"]))
        out.append(("body", trigger(raw["id"], model)))
        out.append(("meta", "Setting: " + ENVIRONMENTS.get(raw.get("environment"), raw.get("environment", "current stage")) + (" | Music: " + raw["music"] if raw.get("music") else " | Music: current stage")))
        for index, original in enumerate(raw["shots"], 1):
            beat_id = original.get("id", f"beat-{index:02d}")
            full_id = raw["id"] + "/" + beat_id
            active = [r for r in ROUTES if not original.get("routes") or r in original["routes"]]
            resolved = {r: dict(original, **original.get("routeDialogue", {}).get(r, {})) for r in active}
            out.append(("beat", f"Beat {index:02d}" + (" " + beat_id.replace('-', ' ') if original.get('id') else "")))
            out.append(("meta", full_id + " | " + timing(original) + (" | Jay route only" if active == ["hero"] else " | Franklin route only" if active == ["franklin"] else "")))
            actions = {r: human_action(resolved[r], r) for r in active}
            if len(set(actions.values())) == 1:
                out.append(("staging", next(iter(actions.values()))))
            else:
                pieces = {r: actions[r].rstrip(".").split("; ") for r in active}
                shared = [p for p in pieces[active[0]] if all(p in pieces[r] for r in active)]
                if shared:
                    out.append(("staging", "; ".join(shared) + "."))
                for r in active:
                    distinct = [p for p in pieces[r] if p not in shared]
                    if distinct:
                        out.append(("staging", ROUTES[r] + ": " + "; ".join(distinct) + "."))
            if len({(s.get("speaker"), s.get("dialogue")) for s in resolved.values()}) == 1:
                choices = [("shared", next(iter(resolved.values())))]
            else:
                choices = list(resolved.items())
            for route_key, shot in choices:
                if shot.get("dialogue"):
                    label = shot.get("speaker", "")
                    if route_key != "shared":
                        label += " / " + ROUTES[route_key]
                    out.append(("speaker", label))
                    out.append(("dialogue", shot["dialogue"]))
                elif shot.get("speaker"):
                    out.append(("meta", shot["speaker"] + " portrait; silent reveal."))
                if shot.get("portrait"):
                    out.append(("meta", "Portrait: " + shot["portrait"] + " | Expression: " + shot.get("expression", "neutral")))
            if not any(s.get("dialogue") for s in resolved.values()):
                out.append(("meta", "No spoken dialogue."))
            for key, label in (("objective", "Objective"), ("caption", "Caption"), ("destination", "Destination")):
                value = original.get(key) or (raw.get(key) if key == "destination" else None)
                if value:
                    out.append(("onscreen", label + ": " + value))
            # An empty separator marks the end of the beat for page grouping.
            out.append(("separator", ""))
    out.extend([("heading1", "Unlock and results text"),
                ("body", "Franklin unlocks when Jay defeats him and leaves Stage 4 without dying during that attempt. Franklin’s own route uses Shermometer v3 in Stage 4. Existing profiles retain their unlock; results do not announce it as a new reward again.")])
    seen = set()
    for entry in model["ui"]:
        key = (entry["id"], entry["text"])
        if key in seen:
            continue
        seen.add(key)
        out.append(("speaker", UI_IDS.get(entry["id"], "Notification")))
        out.append(("dialogue", entry["text"]))
    out.extend([("heading1", "Editing reference"),
                ("body", "Story lines and staging are authored in data/cutscenes.js. Stage and boss triggers are defined in data/campaign.js and engine.js. Unlock, status, and results messages are in app.js."),
                ("body", "Use the scene and beat identifiers when returning revisions. Replace the spoken line under the appropriate route; shared lines affect both characters. The seven approved opening lines are preserved exactly."),
                ("meta", "Source fingerprint: " + model["sourceDigest"][:16])])
    return out


def markdown(block_list: list[tuple[str, str]]) -> str:
    parts = []
    for kind, text in block_list:
        if not text:
            continue
        if kind == "title":
            parts.append("# " + text)
        elif kind in ("heading1", "opening"):
            parts.append("## " + text)
        elif kind == "beat":
            parts.append("### " + text)
        elif kind == "speaker":
            parts.append("**" + text + "**")
        elif kind == "dialogue":
            parts.append(text)
        elif kind == "staging":
            parts.append("*" + text + "*")
        else:
            parts.append(text)
    return "\n\n".join(parts) + "\n"


def word(block_list: list[tuple[str, str]], digest: str):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:" + attr), "Calibri")
        for border in list(style.element.xpath(".//w:pBdr")):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(23)
    doc.styles["Title"].paragraph_format.space_after = Pt(7)
    doc.styles["Subtitle"].font.size = Pt(14)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(16)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Heading 2"].font.size = Pt(11)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(4)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Complete editable Jay and Franklin campaign script"
    doc.core_properties.author = "Cat Town Films"
    doc.core_properties.identifier = digest
    doc.core_properties.keywords = "campaign, dialogue, Jay Sherman, Franklin, editable script"
    beat_paragraphs = []
    scene_paragraphs = []
    for kind, text in block_list:
        if kind == "separator":
            for p in beat_paragraphs[:-1]:
                p.paragraph_format.keep_with_next = True
            if beat_paragraphs:
                beat_paragraphs[-1].paragraph_format.keep_with_next = False
            beat_paragraphs = []
            continue
        style = {"title": "Title", "subtitle": "Subtitle", "heading1": "Heading 1", "opening": "Heading 1", "beat": "Heading 2"}.get(kind, "Normal")
        p = doc.add_paragraph(text, style)
        p.paragraph_format.widow_control = True
        if kind in ("heading1", "opening"):
            scene_paragraphs = [p]
        elif scene_paragraphs and kind != "beat":
            scene_paragraphs.append(p)
        if kind == "opening":
            p.paragraph_format.page_break_before = True
        if kind == "beat":
            for intro in scene_paragraphs:
                intro.paragraph_format.keep_with_next = True
            scene_paragraphs = []
            beat_paragraphs = [p]
        elif beat_paragraphs:
            beat_paragraphs.append(p)
        if kind in ("meta", "staging"):
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.size = Pt(9.5 if kind == "meta" else 10)
                if kind == "staging":
                    r.italic = True
        if kind == "speaker":
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.bold = True
        if kind == "dialogue":
            p.paragraph_format.left_indent = Inches(.15)
            p.paragraph_format.space_after = Pt(5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Script  |  ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--markdown", type=Path, default=ROOT / "docs/FULL-SCRIPT.md")
    parser.add_argument("--docx", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    model = source_model()
    content = blocks(model)
    md = markdown(content)
    report = {"version": model["version"], "sourceDigest": model["sourceDigest"],
              "sourceHashes": model["sourceHashes"], "scenes": len(model["scenes"]),
              "beats": sum(len(s["raw"]["shots"]) for s in model["scenes"]),
              "dialogueEntriesByRoute": {r: sum(bool(shot.get("dialogue")) for s in model["scenes"] for shot in s["routes"][r]["shots"]) for r in ROUTES},
              "activeCampaignDialogueEntriesByRoute": {r: sum(bool(shot.get("dialogue")) for s in model["scenes"] for shot in s["routes"][r]["shots"]) for r in ROUTES},
              "optionalSceneAppendix": []}
    if args.check:
        if not args.markdown.is_file() or args.markdown.read_text() != md:
            raise SystemExit("Script Markdown is stale. Run tools/export_script.py.")
        if args.docx:
            from docx import Document
            actual = Document(args.docx)
            expected = [text for kind, text in content if kind != "separator"]
            if actual.core_properties.identifier != model["sourceDigest"] or [p.text for p in actual.paragraphs] != expected:
                raise SystemExit("Script Word document is stale. Re-export and render before delivery.")
        print(json.dumps({"check": "passed", **report}, indent=2))
        return 0
    args.markdown.parent.mkdir(parents=True, exist_ok=True)
    args.markdown.write_text(md)
    report["markdownSha256"] = sha(args.markdown.read_bytes())
    if args.docx:
        args.docx.parent.mkdir(parents=True, exist_ok=True)
        word(content, model["sourceDigest"]).save(args.docx)
        report["docxSha256"] = sha(args.docx.read_bytes())
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
